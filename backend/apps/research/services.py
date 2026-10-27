import re
import statistics
import uuid
import zipfile
from xml.etree import ElementTree
from pathlib import Path

from django.conf import settings
from rest_framework.exceptions import ValidationError

ALLOWED = {
    '.pdf': {'application/pdf'},
    '.docx': {'application/vnd.openxmlformats-officedocument.wordprocessingml.document'},
    '.txt': {'text/plain', 'application/octet-stream'},
}
EXTRACTION_ALGORITHM = 'academic-docx-v2'
HEADING_ALIASES = {
    'abstract': ('abstract',), 'keywords': ('keywords', 'key words'),
    'introduction': ('introduction', 'background', 'problem statement'),
    'literature_review': ('literature review', 'literature survey', 'related literature'),
    'related_work': ('related work', 'prior work', 'previous work'),
    'research_gap': ('research gap', 'knowledge gap', 'problem gap'),
    'research_objectives': ('research objectives', 'objectives', 'aim and objectives'),
    'methodology': ('methodology', 'methods', 'research methods', 'experimental setup', 'materials and methods', 'proposed method'),
    'dataset': ('dataset description', 'data description', 'dataset'),
    'preprocessing': ('data preprocessing', 'preprocessing'),
    'results': ('results', 'findings', 'experimental results', 'evaluation', 'performance evaluation', 'comparative analysis'),
    'discussion': ('discussion', 'results and discussion'),
    'limitations': ('limitations', 'study limitations'),
    'conclusion': ('conclusion', 'conclusions', 'future work', 'conclusion and future work'),
    'acknowledgments': ('acknowledgments', 'acknowledgements'),
    'references': ('references', 'bibliography', 'works cited', 'reference list'),
    'appendix': ('appendix', 'appendices', 'supplementary material'),
}
KNOWN_TITLES = {alias for aliases in HEADING_ALIASES.values() for alias in aliases} | {
    'data collection', 'feature selection', 'model development', 'statistical analysis',
    'evaluation metrics', 'performance comparison', 'error analysis', 'ethics statement',
}
W_NS = 'http://schemas.openxmlformats.org/wordprocessingml/2006/main'
M_NS = 'http://schemas.openxmlformats.org/officeDocument/2006/math'
XML_NS = {'w': W_NS, 'm': M_NS}


def validate_upload(upload):
    ext = Path(upload.name).suffix.lower()
    if ext not in ALLOWED:
        raise ValidationError({'file': ['Only PDF, DOCX and TXT documents are supported.']})
    if upload.size <= 0 or upload.size > settings.MAX_RESEARCH_UPLOAD_BYTES:
        raise ValidationError({'file': [f'File must be between 1 byte and {settings.MAX_RESEARCH_UPLOAD_BYTES} bytes.']})
    mime = (getattr(upload, 'content_type', '') or 'application/octet-stream').split(';')[0].lower()
    if mime not in ALLOWED[ext]:
        raise ValidationError({'file': ['The file content type does not match its extension.']})
    header = upload.read(8)
    upload.seek(0)
    if ext == '.pdf' and not header.startswith(b'%PDF-'):
        raise ValidationError({'file': ['The uploaded file is not a valid PDF.']})
    if ext == '.docx' and not header.startswith(b'PK'):
        raise ValidationError({'file': ['The uploaded file is not a valid DOCX document.']})
    if ext == '.docx':
        try:
            with zipfile.ZipFile(upload) as archive:
                names = set(archive.namelist())
                if '[Content_Types].xml' not in names or 'word/document.xml' not in names:
                    raise ValidationError({'file': ['The uploaded file is not a valid DOCX document.']})
                if sum(item.file_size for item in archive.infolist()) > settings.MAX_RESEARCH_UPLOAD_BYTES * 20:
                    raise ValidationError({'file': ['The DOCX expands beyond the permitted processing size.']})
            upload.seek(0)
        except zipfile.BadZipFile as exc:
            raise ValidationError({'file': ['The uploaded file is not a valid DOCX document.']}) from exc
    return ext, mime


def save_private_file(upload):
    root = Path(settings.PRIVATE_DOCUMENT_ROOT)
    root.mkdir(parents=True, exist_ok=True)
    key = f'{uuid.uuid4().hex}{Path(upload.name).suffix.lower()}'
    with (root / key).open('wb') as output:
        for chunk in upload.chunks():
            output.write(chunk)
    return key


_NUMBER_PREFIX = re.compile(r'^\s*(?P<number>(?:(?:\d+|[IVXLCDM]+|[A-Z])(?:[.)](?:\d+|[IVXLCDM]+|[A-Z]))*[.)]?|\([A-Za-z0-9ivxlcdm]+\))\s+)(?P<title>.+?)\s*$', re.I)


def _normalize_title(value):
    value = re.sub(r'\s+', ' ', value).strip()
    match = _NUMBER_PREFIX.match(value)
    number = match.group('number').strip() if match else ''
    title = match.group('title') if match else value
    title = re.sub(r'^[\s:.)\-\u2013\u2014]+|[\s:.)\-\u2013\u2014]+$', '', title)
    normalized = re.sub(r'[^a-z0-9 ]+', ' ', title.casefold())
    return ' '.join(normalized.split()), number, title


def _category(title):
    normalized, _, clean = _normalize_title(title)
    for category, aliases in HEADING_ALIASES.items():
        if normalized in aliases:
            return category
    if normalized in KNOWN_TITLES:
        return 'other'
    return None


def _paragraph_xml_text(paragraph):
    # Iterate actual text nodes so text in equations and hyperlinks is not
    # silently lost by python-docx's convenience .text property.
    from docx.oxml.ns import qn
    pieces = []
    for node in paragraph._p.iter():
        if node.tag in (qn('w:t'), f'{{{M_NS}}}t'):
            pieces.append(node.text or '')
        elif node.tag == qn('w:tab'):
            pieces.append('\t')
        elif node.tag in (qn('w:br'), qn('w:cr')):
            pieces.append('\n')
    return ''.join(pieces)


def _paragraph_signals(paragraph, body_font_size):
    text = _paragraph_xml_text(paragraph) if paragraph is not None else ''
    compact = ' '.join(text.split())
    style = getattr(paragraph, 'style', None)
    style_names, style_level = [], None
    seen_styles = set()
    while style is not None and id(style) not in seen_styles:
        seen_styles.add(id(style))
        style_names.append((getattr(style, 'name', '') or '').strip())
        style_match = re.search(r'heading\s*(\d+)', style_names[-1], re.I)
        if style_level is None and style_match:
            style_level = int(style_match.group(1))
        style = getattr(style, 'base_style', None)
    style_name = ' / '.join(style_names)
    direct_properties = paragraph._p.pPr
    style_properties = paragraph.style._element.pPr if paragraph.style is not None else None
    outline = direct_properties.find(f'{{{W_NS}}}outlineLvl') if direct_properties is not None else None
    if outline is None and style_properties is not None:
        outline = style_properties.find(f'{{{W_NS}}}outlineLvl')
    outline_level = int(outline.get(f'{{{W_NS}}}val')) + 1 if outline is not None else None
    numbering = direct_properties.find(f'{{{W_NS}}}numPr') if direct_properties is not None else None
    if numbering is None and style_properties is not None:
        numbering = style_properties.find(f'{{{W_NS}}}numPr')
    ilvl = numbering.find(f'{{{W_NS}}}ilvl') if numbering is not None else None
    numid = numbering.find(f'{{{W_NS}}}numId') if numbering is not None else None
    number, _, _ = _normalize_title(compact)
    numeric_prefix = _NUMBER_PREFIX.match(compact)
    runs = [run for run in paragraph.runs if run.text.strip()]
    bold_ratio = sum(bool(run.bold) for run in runs) / len(runs) if runs else 0
    sizes = [run.font.size.pt for run in runs if run.font.size]
    relative_large = bool(body_font_size and sizes and statistics.median(sizes) >= body_font_size * 1.12)
    alignment = str(paragraph.alignment).lower() if paragraph.alignment is not None else ''
    centered = alignment.endswith('center')
    all_caps = any(char.isalpha() for char in compact) and compact.upper() == compact
    known_category = _category(compact)
    heading_style = any('heading' in name.casefold() or 'section' in name.casefold() for name in style_names)
    style_level = style_level if style_level is not None else outline_level
    if style_level is None and ilvl is not None:
        try:
            style_level = int(ilvl.get(f'{{{W_NS}}}val')) + 1
        except (TypeError, ValueError):
            pass
    if style_level is None and numeric_prefix:
        components = re.findall(r'\d+|[IVXLCDM]+|[A-Z]', numeric_prefix.group('number').rstrip('.)'), re.I)
        style_level = max(1, len(components), 2 if len(components) == 1 and len(components[0]) == 1 and components[0].isalpha() else 1)
    short = 1 <= len(compact) <= 140
    sentence_like = bool(re.search(r'[.!?][\]\)\"\']?$', compact))
    academic_known = known_category is not None
    # Canonical academic names are strong semantic evidence. Unrecognized
    # headings require explicit structural/formatting evidence and a title-like
    # shape; bold alone is intentionally insufficient.
    structural = heading_style or outline_level is not None
    formatted = bold_ratio >= 0.8 and (centered or all_caps or relative_large or bool(numeric_prefix))
    recognized = short and (academic_known or (structural and not sentence_like) or
                            (formatted and not sentence_like and (bool(numeric_prefix) or centered or all_caps)))
    if not recognized:
        return None
    if style_level is None:
        style_level = 1
    signals = []
    if academic_known: signals.append('academic_title')
    if heading_style: signals.append('heading_style')
    if outline_level is not None: signals.append('outline_level')
    if numbering is not None: signals.append('word_numbering')
    if numeric_prefix: signals.append('text_numbering')
    if bold_ratio >= 0.8: signals.append('bold')
    if centered: signals.append('centered')
    if all_caps: signals.append('uppercase')
    if relative_large: signals.append('larger_font')
    title = _normalize_title(compact)[2]
    confidence = 'high' if academic_known or structural else 'medium'
    return {'title': compact, 'normalized_title': title, 'category': known_category or 'other',
            'numbering': numeric_prefix.group('number').strip() if numeric_prefix else None,
            'level': max(1, min(6, style_level)), 'signals': signals,
            'confidence': confidence, 'num_id': numid.get(f'{{{W_NS}}}val') if numid is not None else None,
            'numbering_level': ilvl.get(f'{{{W_NS}}}val') if ilvl is not None else None}


def _plain_text_heading(text):
    compact = ' '.join(text.split())
    normalized, number, title = _normalize_title(compact)
    category = _category(compact)
    if not category or len(compact) > 140 or re.search(r'[.!?][\]\)\"\']?$', compact):
        return None
    return {'title': compact, 'normalized_title': title, 'category': category,
            'numbering': number or None, 'level': 1, 'signals': ['academic_title'] if not number else ['academic_title', 'text_numbering'],
            'confidence': 'high', 'num_id': None, 'numbering_level': None}


def _body_font_size(doc):
    sizes = []
    for paragraph in doc.paragraphs:
        for run in paragraph.runs:
            if run.text.strip() and run.font.size:
                sizes.append(run.font.size.pt)
    return statistics.median(sizes) if sizes else None


def _docx_blocks(doc):
    from docx.text.paragraph import Paragraph
    from docx.table import Table
    body = doc._element.body
    blocks, tables, order, unsupported = [], [], 0, []
    for element in body.iterchildren():
        if element.tag == f'{{{W_NS}}}p':
            paragraph = Paragraph(element, doc)
            text = _paragraph_xml_text(paragraph)
            if text.strip():
                order += 1
                style = getattr(getattr(paragraph, 'style', None), 'name', '') or ''
                blocks.append({'text': text, 'order': order, 'kind': 'paragraph', 'style': style,
                               'paragraph': paragraph, 'source_id': f'p-{order:06d}'})
        elif element.tag == f'{{{W_NS}}}tbl':
            table = Table(element, doc)
            if any(child is not element and child.tag == f'{{{W_NS}}}tbl' for child in element.iter()):
                unsupported.append('nested_table')
            table_order = order + 1
            table_rows, cell_seen = [], set()
            for row_index, row in enumerate(table.rows, start=1):
                row_values = []
                for cell_index, cell in enumerate(row.cells, start=1):
                    tc_id = id(cell._tc)
                    cell_records = []
                    if tc_id not in cell_seen:
                        cell_seen.add(tc_id)
                        for paragraph_index, paragraph in enumerate(cell.paragraphs, start=1):
                            text = _paragraph_xml_text(paragraph)
                            if text.strip():
                                order += 1
                                source_id = f't-{table_order:05d}-r{row_index:03d}-c{cell_index:03d}-p{paragraph_index:03d}'
                                style = getattr(getattr(paragraph, 'style', None), 'name', '') or ''
                                record = {'text': text, 'order': order, 'kind': 'table_paragraph', 'style': style,
                                          'paragraph': paragraph, 'source_id': source_id,
                                          'table_order': table_order, 'row': row_index, 'cell': cell_index}
                                blocks.append(record)
                                cell_records.append(record)
                    row_values.append(cell_records)
                table_rows.append(row_values)
            tables.append({'table_order': table_order, 'rows': table_rows})
            order += 1
    for element in doc._element.iter():
        if element.tag.endswith('}sdt'):
            unsupported.append('content_control')
        elif element.tag.endswith('}txbxContent'):
            unsupported.append('text_box')
    for section in doc.sections:
        if any(node.tag in (f'{{{W_NS}}}t', f'{{{M_NS}}}t') and (node.text or '').strip()
               for part in (section.header._element, section.footer._element) for node in part.iter()):
            unsupported.append('header_or_footer_text')
    return blocks, tables, unsupported


def _notes(path):
    results, unsupported = [], []
    with zipfile.ZipFile(path) as archive:
        for part, kind in (('word/footnotes.xml', 'footnote'), ('word/endnotes.xml', 'endnote')):
            if part not in archive.namelist():
                continue
            try:
                from lxml import etree
                root = etree.fromstring(archive.read(part))
                for note in root:
                    note_id = note.get(f'{{{W_NS}}}id')
                    if note_id in ('-1', '0'):
                        continue
                    text = ''.join(note.xpath('.//w:t/text() | .//m:t/text()', namespaces=XML_NS))
                    if text.strip():
                        results.append({'kind': kind, 'id': note_id, 'text': text})
            except Exception:
                unsupported.append(kind)
    return results, unsupported


def _docx_source_text_block_count(path):
    with zipfile.ZipFile(path) as archive:
        root = ElementTree.fromstring(archive.read('word/document.xml'))
    count = 0
    for paragraph in root.iter(f'{{{W_NS}}}p'):
        text = ''.join((node.text or '') for node in paragraph.iter()
                       if node.tag in (f'{{{W_NS}}}t', f'{{{M_NS}}}t'))
        if text.strip():
            count += 1
    return count


def _build_sections(blocks, table_specs, pages, metadata, extra_warnings=None, notes=None):
    headings, sections, unclassified = [], [], []
    body_font_size = metadata.pop('_body_font_size', None)
    source_text_count = metadata.pop('_source_text_blocks', len(blocks))
    table_by_order = {item['table_order']: item for item in table_specs}
    section_stack = []
    current = None
    assignments = {}
    character_offset = 0
    for document_index, block in enumerate(blocks):
        start_offset = character_offset
        end_offset = start_offset + len(block['text'])
        character_offset = end_offset + 1
        detected = (_paragraph_signals(block.get('paragraph'), body_font_size)
                    if block.get('paragraph') is not None else _plain_text_heading(block['text']))
        if detected:
            level = detected['level']
            while section_stack and section_stack[-1]['level'] >= level:
                section_stack.pop()
            parent = section_stack[-1] if section_stack else None
            section = {'title': detected['title'], 'original_title': block['text'], 'normalized_title': detected['normalized_title'],
                       'category': detected['category'], 'numbering': detected['numbering'],
                       'level': level, 'parent_key': parent['section_key'] if parent else None,
                       'document_order': block['order'], 'paragraphs': [], 'paragraph_records': [],
                       'content_blocks': [], 'heading_source_id': block['source_id'],
                       'confidence': detected['confidence'], 'signals': detected['signals'],
                       'section_key': f"{len(sections)}-{uuid.uuid5(uuid.NAMESPACE_URL, block['source_id'] + '|' + detected['normalized_title']).hex[:16]}"}
            sections.append(section)
            section_stack.append(section)
            current = section
            headings.append({'title': section['title'], 'original_text': block['text'], 'normalized_category': section['category'],
                             'numbering': section['numbering'], 'level': level, 'parent_key': section['parent_key'],
                             'paragraph_index': document_index, 'source_id': block['source_id'],
                             'start_offset': start_offset, 'end_offset': end_offset,
                             'confidence': section['confidence'], 'signals': section['signals'],
                             'word_numbering_id': detected['num_id'], 'word_numbering_level': detected['numbering_level']})
            assignments[block['source_id']] = 'heading'
            continue

        record = {'source_id': block['source_id'], 'document_order': block['order'], 'text': block['text'],
                  'start_offset': start_offset, 'end_offset': end_offset,
                  'kind': block['kind'], 'paragraph_index': len(current['paragraphs']) if current else None,
                  'table_order': block.get('table_order'), 'row': block.get('row'), 'cell': block.get('cell')}
        if current:
            current['paragraphs'].append(block['text'])
            current['paragraph_records'].append(record)
            assignments[block['source_id']] = current['section_key']
            if block['kind'] == 'table_paragraph':
                table = table_by_order.get(block['table_order'])
                if table:
                    content_block = next((item for item in current['content_blocks'] if item.get('table_order') == block['table_order']), None)
                    if content_block is None:
                        content_block = {'type': 'table', 'table_order': block['table_order'], 'rows': []}
                        current['content_blocks'].append(content_block)
                    while len(content_block['rows']) < block['row']:
                        content_block['rows'].append([])
                    row = content_block['rows'][block['row'] - 1]
                    while len(row) < block['cell']:
                        row.append([])
                    row[block['cell'] - 1].append(record)
            else:
                current['content_blocks'].append({'type': 'paragraph', 'record': record})
        else:
            unclassified.append(record)
            assignments[block['source_id']] = 'unclassified'

    # Make all supported prose available to analysis without hiding pre-heading
    # material, while retaining an explicit unclassified record for reviewers.
    if unclassified:
        section_key = f"{len(sections)}-{uuid.uuid5(uuid.NAMESPACE_URL, 'unclassified|' + str(unclassified[0]['source_id'])).hex[:16]}"
        for index, record in enumerate(unclassified):
            record['paragraph_index'] = index
        sections.insert(0, {'title': 'Unclassified content', 'normalized_title': 'unclassified content',
                             'category': 'other', 'numbering': None, 'level': 1, 'parent_key': None,
                             'document_order': unclassified[0]['document_order'], 'paragraphs': [r['text'] for r in unclassified],
                             'paragraph_records': unclassified, 'content_blocks': [{'type': 'paragraph', 'record': r} for r in unclassified],
                             'section_key': section_key, 'confidence': 'none', 'signals': [], 'unclassified': True})
    # Preserve the existing integer source index contract and compute unique
    # assignment/coverage metrics over every non-empty source text block.
    flat = [item['text'] for item in blocks]
    ids = [item['source_id'] for item in blocks]
    duplicate_assignments = len(ids) - len(set(ids))
    text_count = len(blocks)
    assigned_count = sum(1 for source_id in ids if source_id in assignments)
    warnings = list(extra_warnings or [])
    empty_sections = [s['title'] for s in sections if not s['paragraphs']]
    if empty_sections:
        warnings.append('Some detected headings have no following text before the next heading.')
    if unclassified:
        warnings.append('Text before the first confidently detected heading is retained as unclassified content.')
    if duplicate_assignments:
        warnings.append('Duplicate source block identifiers were detected; inspect the document structure.')
    if not headings:
        warnings.append('No confident academic headings were detected; all text is retained as unclassified content.')
    unsupported = sum(int(match.group(1)) if (match := re.search(r'\((\d+) occurrence', warning)) else 1
                      for warning in (extra_warnings or []) if warning.startswith('Unsupported element:'))
    coverage = {'source_text_blocks': source_text_count, 'extracted_text_blocks': text_count,
                'detected_headings': len(headings), 'top_level_sections': sum(h['level'] == 1 for h in headings),
                'subsections': sum(h['level'] > 1 for h in headings), 'unclassified_blocks': len(unclassified),
                'skipped_unsupported_elements': unsupported, 'duplicate_assignments': duplicate_assignments,
                'assigned_blocks': assigned_count,
                'coverage_percent': (round(100 * assigned_count / source_text_count, 2) if source_text_count else 100.0) if not unsupported else None,
                'empty_sections': empty_sections}
    reference = next((s for s in sections if s['category'] == 'references'), None)
    return {'extraction_algorithm': EXTRACTION_ALGORITHM, 'pages': pages, 'paragraphs': flat,
            'headings': headings, 'sections': sections, 'unclassified': unclassified,
            'tables': [{'table_order': t['table_order'], 'rows': [[''.join(p['text'] for p in cell) for cell in row] for row in t['rows']]} for t in table_specs],
            'notes': notes or [], 'coverage': {**coverage, 'extracted_notes': len(notes or []),
                                               'scope': 'non-empty ordered main-body text blocks; notes are reported separately'},
            'warnings': list(dict.fromkeys(warnings)),
            'reference_section': reference, 'metadata': metadata}


def _pdf_paragraphs(text):
    paragraphs = []
    for block in re.split(r'\n\s*\n', text):
        for line in (line.strip() for line in block.splitlines() if line.strip()):
            if _plain_text_heading(line) or not paragraphs or _plain_text_heading(paragraphs[-1]):
                paragraphs.append(line)
            else:
                paragraphs[-1] = f'{paragraphs[-1]} {line}'
    return paragraphs


def extract_document(path, extension):
    pages, paragraphs = [], []
    extra_warnings = []
    if extension == '.pdf':
        import pymupdf
        with pymupdf.open(path) as pdf:
            for index, page in enumerate(pdf):
                text = page.get_text('text')
                pages.append({'page': index + 1, 'text': text})
                paragraphs.extend(_pdf_paragraphs(text))
            metadata = {k: v for k, v in (pdf.metadata or {}).items() if v}
        blocks = [{'text': text, 'order': i, 'kind': 'paragraph', 'source_id': f'p-{i:06d}'} for i, text in enumerate(paragraphs, 1)]
        table_specs = []
    elif extension == '.docx':
        from docx import Document
        doc = Document(path)
        blocks, table_specs, unsupported_elements = _docx_blocks(doc)
        paragraphs = [item['text'] for item in blocks]
        properties = doc.core_properties
        metadata = {key: value for key, value in {
            'title': properties.title, 'author': properties.author, 'subject': properties.subject,
            'paragraph_count': sum(item['kind'] == 'paragraph' for item in blocks),
            'table_paragraph_count': sum(item['kind'] == 'table_paragraph' for item in blocks),
        }.items() if value}
        metadata['_body_font_size'] = _body_font_size(doc)
        metadata['_source_text_blocks'] = _docx_source_text_block_count(path)
        notes, unsupported_notes = _notes(path)
        if unsupported_notes:
            extra_warnings.append('Unsupported element: footnote/endnote XML could not be parsed.')
        if notes:
            extra_warnings.append('Footnotes/endnotes were extracted separately; exact inline anchor positions are not mapped.')
        for element_name in set(unsupported_elements):
            extra_warnings.append(f'Unsupported element: {element_name} ({unsupported_elements.count(element_name)} occurrence(s)); some text may not be included in ordered body extraction.')
        if any('m:oMath' in element.tag for element in doc._element.iter()):
            extra_warnings.append('Equation text nodes were retained where readable; visual equation layout is not represented.')
    else:
        raw = Path(path).read_bytes()
        try:
            text = raw.decode('utf-8-sig')
        except UnicodeDecodeError:
            try:
                text = raw.decode('utf-16')
            except UnicodeDecodeError as exc:
                raise ValueError('TXT must use UTF-8 or UTF-16 encoding.') from exc
        paragraphs = [part for part in re.split(r'\n\s*\n', text) if part.strip()]
        metadata = {'character_count': len(text)}
        blocks = [{'text': text, 'order': i, 'kind': 'paragraph', 'source_id': f'p-{i:06d}'} for i, text in enumerate(paragraphs, 1)]
        table_specs, notes = [], []
    metadata.setdefault('paragraph_count', len(blocks))
    return _build_sections(blocks, table_specs, pages, metadata, extra_warnings, locals().get('notes', []))
