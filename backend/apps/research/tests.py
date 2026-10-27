import io
import tempfile
from datetime import timedelta
from pathlib import Path
from unittest.mock import patch

import pymupdf
from django.contrib.auth import get_user_model
from django.core.files.uploadedfile import SimpleUploadedFile
from django.test import TestCase, override_settings
from django.utils import timezone
from docx import Document
from rest_framework.test import APIClient

from .models import DocumentVersion, ResearchDocument, ResearchProject
from .services import EXTRACTION_ALGORITHM, extract_document
from .tasks import process_document_version

User = get_user_model()


class ResearchApiTests(TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.settings = override_settings(PRIVATE_DOCUMENT_ROOT=Path(self.temp.name), MAX_RESEARCH_UPLOAD_BYTES=2_000_000,
                                         CELERY_TASK_ALWAYS_EAGER=True)
        self.settings.enable()
        self.user = User.objects.create_user(username='researcher-one', email='one@example.test', password='ValidPassword!42', is_active=True)
        self.other = User.objects.create_user(username='researcher-two', email='two@example.test', password='ValidPassword!42', is_active=True)
        self.client = APIClient()
        self.client.force_authenticate(self.user)
        self.project_url = '/api/v1/projects/'

    def tearDown(self):
        self.settings.disable()
        self.temp.cleanup()

    def project(self, title='Field study'):
        return ResearchProject.objects.create(owner=self.user, title=title)

    def pdf(self, text='Abstract\n\nOriginal abstract.\n\nIntroduction\n\nStudy text.'):
        pdf = pymupdf.open()
        page = pdf.new_page()
        page.insert_text((72, 72), text)
        raw = pdf.tobytes()
        pdf.close()
        return SimpleUploadedFile('study.pdf', raw, content_type='application/pdf')

    def test_project_crud_and_owner_isolation(self):
        response = self.client.post(self.project_url, {'title': 'Water study', 'keywords': ['water']}, format='json')
        self.assertEqual(response.status_code, 201)
        project_id = response.data['id']
        self.assertEqual(self.client.get(self.project_url).data[0]['title'], 'Water study')
        self.assertEqual(self.client.patch(f'{self.project_url}{project_id}/', {'title': 'Edited'}, format='json').status_code, 200)
        self.client.force_authenticate(self.other)
        self.assertEqual(self.client.get(f'{self.project_url}{project_id}/').status_code, 404)
        self.assertEqual(self.client.delete(f'{self.project_url}{project_id}/').status_code, 404)

    def test_pdf_upload_processing_and_versions(self):
        project = self.project()
        response = self.client.post(f'/api/v1/projects/{project.pk}/documents/upload/', {'file': self.pdf()}, format='multipart')
        self.assertEqual(response.status_code, 201)
        document = ResearchDocument.objects.get(pk=response.data['id'])
        version = document.current_version
        self.assertEqual(version.processing_status, DocumentVersion.ProcessingStatus.COMPLETED)
        self.assertIn('Abstract', [section['title'] for section in version.extraction_metadata['sections']])
        self.assertTrue((Path(self.temp.name) / version.storage_key).exists())
        second = self.client.post(f'/api/v1/projects/{project.pk}/documents/upload/', {'file': self.pdf(), 'document_id': str(document.pk)}, format='multipart')
        self.assertEqual(second.status_code, 201)
        self.assertEqual(document.versions.count(), 2)
        self.assertNotIn('storage_key', second.data['current_version'])
        destination = self.project('Second project')
        moved = self.client.patch(f'/api/v1/documents/{document.pk}/', {'project': str(destination.pk)}, format='json')
        self.assertEqual(moved.status_code, 200)
        self.assertEqual(moved.data['project'], destination.pk)

    def test_login_project_upload_extraction_and_private_document_view(self):
        login = self.client.post('/api/v1/auth/login/', {'username': self.user.email, 'password': 'ValidPassword!42'}, format='json')
        self.assertEqual(login.status_code, 200)
        self.client.credentials(HTTP_AUTHORIZATION=f"Bearer {login.data['access']}")
        project = self.client.post(self.project_url, {'title': 'Workflow project'}, format='json')
        self.assertEqual(project.status_code, 201)
        uploaded = self.client.post(f"/api/v1/projects/{project.data['id']}/documents/upload/", {'file': self.pdf()}, format='multipart')
        self.assertEqual(uploaded.status_code, 201)
        document_id = uploaded.data['id']
        status_response = self.client.get(f'/api/v1/documents/{document_id}/status/')
        self.assertEqual(status_response.data['status'], DocumentVersion.ProcessingStatus.COMPLETED)
        self.assertIn('Introduction', [s['title'] for s in status_response.data['extraction_metadata']['sections']])
        file_response = self.client.get(f'/api/v1/documents/{document_id}/file/')
        self.assertEqual(file_response.status_code, 200)
        self.assertEqual(file_response['Content-Type'], 'application/pdf')
        file_response.close()

    def test_docx_and_txt_extraction_preserve_paragraphs(self):
        doc = Document()
        doc.add_paragraph('Methodology')
        doc.add_paragraph('Collected original evidence.')
        stream = io.BytesIO()
        doc.save(stream)
        path = Path(self.temp.name) / 'sample.docx'
        path.write_bytes(stream.getvalue())
        result = extract_document(path, '.docx')
        self.assertEqual(result['sections'][0]['title'], 'Methodology')
        self.assertEqual(result['sections'][0]['paragraphs'], ['Collected original evidence.'])
        txt = Path(self.temp.name) / 'sample.txt'
        txt.write_text('Results\n\nRaw result line.', encoding='utf-8')
        self.assertEqual(extract_document(txt, '.txt')['sections'][0]['paragraphs'], ['Raw result line.'])

    def test_docx_styled_hierarchy_tables_captions_and_coverage(self):
        doc = Document()
        doc.add_paragraph('Front matter is retained, not discarded.')
        doc.add_paragraph('Abstract', style='Heading 1')
        doc.add_paragraph('Only the abstract belongs in this section.')
        doc.add_paragraph('III. METHODOLOGY', style='Heading 1')
        custom = doc.styles.add_style('Custom Subsection', 1)
        custom.base_style = doc.styles['Heading 2']
        doc.add_paragraph('A. Dataset Description', style=custom)
        doc.add_paragraph('Dataset source remains exact.')
        table = doc.add_table(rows=2, cols=2)
        table.cell(0, 0).text = 'Model'
        table.cell(0, 1).text = 'Accuracy'
        table.cell(1, 0).text = 'Baseline'
        table.cell(1, 1).text = '0.81'
        doc.add_paragraph('Table 1. Model comparison', style='Caption')
        doc.add_paragraph('IV. RESULTS AND DISCUSSION', style='Heading 1')
        doc.add_paragraph('Comparison results remain in the results section.')
        doc.add_paragraph('References:', style='Heading 1')
        doc.add_paragraph('An actual citation remains unchanged.')
        stream = io.BytesIO(); doc.save(stream)
        path = Path(self.temp.name) / 'structured.docx'; path.write_bytes(stream.getvalue())

        result = extract_document(path, '.docx')
        sections = result['sections']
        by_title = {section['title']: section for section in sections}
        self.assertEqual(by_title['Abstract']['paragraphs'], ['Only the abstract belongs in this section.'])
        self.assertEqual(by_title['III. METHODOLOGY']['level'], 1)
        self.assertEqual(by_title['A. Dataset Description']['level'], 2)
        self.assertEqual(by_title['A. Dataset Description']['parent_key'], by_title['III. METHODOLOGY']['section_key'])
        self.assertTrue(any('Model' in paragraph for paragraph in by_title['A. Dataset Description']['paragraphs']))
        self.assertEqual(by_title['IV. RESULTS AND DISCUSSION']['category'], 'discussion')
        self.assertIn('Table 1. Model comparison', by_title['A. Dataset Description']['paragraphs'])
        self.assertEqual(by_title['References:']['category'], 'references')
        self.assertEqual(result['coverage']['coverage_percent'], 100.0)
        self.assertEqual(result['coverage']['duplicate_assignments'], 0)
        self.assertEqual(len(result['paragraphs']), result['coverage']['source_text_blocks'])
        self.assertTrue(result['unclassified'])
        self.assertTrue(result['tables'][0]['rows'][1][1] == '0.81')
        self.assertTrue(all(record['source_id'] for section in sections for record in section['paragraph_records']))

    def test_docx_manually_formatted_numbered_headings_and_duplicate_text(self):
        doc = Document()
        heading = doc.add_paragraph()
        run = heading.add_run('1. INTRODUCTION'); run.bold = True
        paragraph = doc.add_paragraph('Repeated source sentence.')
        second = doc.add_paragraph()
        run = second.add_run('1.1 Research Objectives'); run.bold = True
        doc.add_paragraph('Repeated source sentence.')
        path = Path(self.temp.name) / 'manual.docx'; doc.save(path)
        result = extract_document(path, '.docx')
        self.assertEqual([h['title'] for h in result['headings']], ['1. INTRODUCTION', '1.1 Research Objectives'])
        self.assertEqual(result['headings'][1]['level'], 2)
        self.assertEqual(result['sections'][0]['paragraphs'], ['Repeated source sentence.'])
        self.assertEqual(result['sections'][1]['paragraphs'], ['Repeated source sentence.'])
        self.assertNotEqual(result['sections'][0]['paragraph_records'][0]['source_id'], result['sections'][1]['paragraph_records'][0]['source_id'])

    def test_docx_without_headings_retains_all_text_as_unclassified(self):
        doc = Document()
        doc.add_paragraph('A plain paragraph.')
        doc.add_paragraph('Another plain paragraph.')
        path = Path(self.temp.name) / 'plain.docx'; doc.save(path)
        result = extract_document(path, '.docx')
        self.assertEqual(result['headings'], [])
        self.assertEqual(result['sections'][0]['title'], 'Unclassified content')
        self.assertEqual(result['sections'][0]['paragraphs'], ['A plain paragraph.', 'Another plain paragraph.'])
        self.assertEqual(result['coverage']['coverage_percent'], 100.0)

    def test_unsupported_nested_table_is_reported_and_coverage_is_not_claimed_complete(self):
        doc = Document()
        doc.add_paragraph('Introduction', style='Heading 1')
        outer = doc.add_table(rows=1, cols=1)
        nested = outer.cell(0, 0).add_table(rows=1, cols=1)
        nested.cell(0, 0).text = 'Nested table text'
        path = Path(self.temp.name) / 'nested.docx'; doc.save(path)
        result = extract_document(path, '.docx')
        self.assertGreater(result['coverage']['skipped_unsupported_elements'], 0)
        self.assertIsNone(result['coverage']['coverage_percent'])
        self.assertTrue(any('nested_table' in warning for warning in result['warnings']))

    def test_reprocess_is_a_new_idempotent_version_and_preserves_source(self):
        project = self.project()
        source_doc = Document(); source_doc.add_paragraph('Introduction', style='Heading 1'); source_doc.add_paragraph('Original text.')
        path = Path(self.temp.name) / 'legacy.docx'; source_doc.save(path)
        from .services import save_private_file
        with path.open('rb') as stream:
            upload = SimpleUploadedFile('legacy.docx', stream.read(), content_type='application/vnd.openxmlformats-officedocument.wordprocessingml.document')
        source_key = save_private_file(upload)
        document = ResearchDocument.objects.create(project=project, owner=self.user, original_filename='legacy.docx')
        source = DocumentVersion.objects.create(document=document, revision_number=1, original_filename='legacy.docx',
            file_extension='.docx', mime_type='application/vnd.openxmlformats-officedocument.wordprocessingml.document',
            file_size=path.stat().st_size, storage_key=source_key, extraction_algorithm='legacy-v1',
            processing_status=DocumentVersion.ProcessingStatus.COMPLETED, extraction_metadata={'sections': [{'title': 'Old', 'paragraphs': ['legacy text']} ]})
        original_bytes = (Path(self.temp.name) / source.storage_key).read_bytes()
        url = f'/api/v1/documents/{document.pk}/versions/{source.pk}/reprocess/'
        first = self.client.post(url, {}, format='json')
        self.assertEqual(first.status_code, 202)
        candidate = DocumentVersion.objects.get(pk=first.data['id'])
        self.assertEqual(candidate.processing_status, DocumentVersion.ProcessingStatus.COMPLETED)
        self.assertEqual(candidate.extraction_algorithm, EXTRACTION_ALGORITHM)
        self.assertEqual(candidate.reprocessed_from_id, source.pk)
        self.assertEqual((Path(self.temp.name) / candidate.storage_key).read_bytes(), original_bytes)
        self.assertEqual(source.extraction_metadata['sections'][0]['title'], 'Old')
        second = self.client.post(url, {}, format='json')
        self.assertEqual(second.status_code, 200)
        self.assertEqual(second.data['id'], first.data['id'])
        self.assertEqual(document.versions.count(), 2)

        self.client.force_authenticate(self.other)
        self.assertEqual(self.client.post(url, {}, format='json').status_code, 404)

    def test_completed_current_version_can_be_reprocessed_into_a_new_revision(self):
        project = self.project()
        source_doc = Document(); source_doc.add_paragraph('Introduction', style='Heading 1'); source_doc.add_paragraph('Original text.')
        path = Path(self.temp.name) / 'current.docx'; source_doc.save(path)
        from .services import save_private_file
        with path.open('rb') as stream:
            upload = SimpleUploadedFile('current.docx', stream.read(), content_type='application/vnd.openxmlformats-officedocument.wordprocessingml.document')
        source_key = save_private_file(upload)
        document = ResearchDocument.objects.create(project=project, owner=self.user, original_filename='current.docx')
        source = DocumentVersion.objects.create(document=document, revision_number=1, original_filename='current.docx',
            file_extension='.docx', mime_type='application/vnd.openxmlformats-officedocument.wordprocessingml.document',
            file_size=path.stat().st_size, storage_key=source_key, extraction_algorithm=EXTRACTION_ALGORITHM,
            processing_status=DocumentVersion.ProcessingStatus.COMPLETED, extraction_metadata={'sections': [{'title': 'Introduction'}]})
        original_bytes = (Path(self.temp.name) / source.storage_key).read_bytes()
        url = f'/api/v1/documents/{document.pk}/versions/{source.pk}/reprocess/'

        response = self.client.post(url, {}, format='json')

        self.assertEqual(response.status_code, 202)
        revision = DocumentVersion.objects.get(pk=response.data['id'])
        self.assertEqual(revision.revision_number, 2)
        self.assertEqual(revision.reprocessed_from_id, source.pk)
        self.assertEqual(revision.extraction_algorithm, EXTRACTION_ALGORITHM)
        self.assertEqual((Path(self.temp.name) / source.storage_key).read_bytes(), original_bytes)
        self.assertEqual(source.extraction_metadata['sections'][0]['title'], 'Introduction')
        self.assertEqual(document.versions.count(), 2)

    def test_invalid_extension_and_size_rejected(self):
        project = self.project()
        wrong = SimpleUploadedFile('bad.exe', b'no', content_type='application/octet-stream')
        self.assertEqual(self.client.post(f'/api/v1/projects/{project.pk}/documents/upload/', {'file': wrong}, format='multipart').status_code, 400)
        with override_settings(MAX_RESEARCH_UPLOAD_BYTES=3):
            self.assertEqual(self.client.post(f'/api/v1/projects/{project.pk}/documents/upload/', {'file': self.pdf()}, format='multipart').status_code, 400)

    def test_unauthenticated_access_and_delete_private_files(self):
        project = self.project()
        result = self.client.post(f'/api/v1/projects/{project.pk}/documents/upload/', {'file': self.pdf()}, format='multipart')
        document = ResearchDocument.objects.get(pk=result.data['id'])
        private = Path(self.temp.name) / document.current_version.storage_key
        self.client.force_authenticate(self.other)
        self.assertEqual(self.client.get(f'/api/v1/documents/{document.pk}/').status_code, 404)
        self.assertEqual(self.client.delete(f'/api/v1/documents/{document.pk}/').status_code, 404)
        self.client.force_authenticate(self.user)
        with self.captureOnCommitCallbacks(execute=True):
            self.assertEqual(self.client.delete(f'/api/v1/documents/{document.pk}/').status_code, 204)
        self.assertFalse(private.exists())
        self.client.force_authenticate(None)
        self.assertEqual(self.client.get(self.project_url).status_code, 401)

    def test_task_failure_is_recorded_without_exposing_error_details(self):
        project = self.project()
        with patch('apps.research.tasks.extract_document', side_effect=ValueError('sensitive parser internals')):
            with override_settings(CELERY_TASK_ALWAYS_EAGER=True, CELERY_TASK_EAGER_PROPAGATES=False):
                result = self.client.post(f'/api/v1/projects/{project.pk}/documents/upload/', {'file': self.pdf()}, format='multipart')
        self.assertEqual(result.status_code, 201)
        version = ResearchDocument.objects.get(pk=result.data['id']).current_version
        self.assertEqual(version.processing_status, DocumentVersion.ProcessingStatus.FAILED)
        self.assertNotIn('sensitive parser internals', str(version.extraction_metadata))

    def test_broker_outage_keeps_upload_pending_and_retry_recovers_same_version_idempotently(self):
        project = self.project()
        with patch('apps.research.views.process_document_version.delay', side_effect=RuntimeError('broker unavailable')):
            uploaded = self.client.post(f'/api/v1/projects/{project.pk}/documents/upload/',
                {'file': self.pdf()}, format='multipart')
        self.assertEqual(uploaded.status_code, 201)
        document = ResearchDocument.objects.get(pk=uploaded.data['id'])
        version = document.current_version
        self.assertEqual(version.processing_status, DocumentVersion.ProcessingStatus.PENDING)
        self.assertIn('queue is unavailable', version.processing_error)
        status_response = self.client.get(f'/api/v1/documents/{document.pk}/status/')
        self.assertEqual(status_response.status_code, 200)
        self.assertIn('queue is unavailable', status_response.data['processing_error'])
        original_key = version.storage_key
        original_version_id = version.pk

        retry_url = f'/api/v1/documents/{document.pk}/versions/{version.pk}/retry-extraction/'
        with patch('apps.research.views.process_document_version.delay') as enqueue:
            queued = self.client.post(retry_url, {}, format='json')
        self.assertEqual(queued.status_code, 202)
        enqueue.assert_called_once_with(str(version.pk))
        version.refresh_from_db()
        self.assertEqual(version.processing_status, DocumentVersion.ProcessingStatus.PENDING)
        self.assertEqual(version.pk, original_version_id)
        self.assertEqual(document.versions.count(), 1)

        def extract_and_assert_claimed(path, extension):
            version.refresh_from_db()
            self.assertEqual(version.processing_status, DocumentVersion.ProcessingStatus.PROCESSING)
            self.assertIsNotNone(version.processing_started_at)
            return {'sections': [{'title': 'Abstract', 'paragraphs': ['Original extracted text.']}]}

        with patch('apps.research.tasks.extract_document', side_effect=extract_and_assert_claimed) as extract:
            result = process_document_version.run(str(version.pk))
            duplicate = process_document_version.run(str(version.pk))
        version.refresh_from_db()
        self.assertEqual(result['status'], DocumentVersion.ProcessingStatus.COMPLETED)
        self.assertTrue(duplicate['duplicate_delivery'])
        self.assertEqual(version.processing_status, DocumentVersion.ProcessingStatus.COMPLETED)
        self.assertEqual(version.processing_attempt, 1)
        self.assertEqual(extract.call_count, 1)
        self.assertEqual(version.storage_key, original_key)
        self.assertTrue((Path(self.temp.name) / original_key).is_file())
        self.assertEqual(document.versions.count(), 1)

    def test_retry_is_owner_scoped_and_refuses_a_fresh_processing_claim(self):
        project = self.project()
        uploaded = self.client.post(f'/api/v1/projects/{project.pk}/documents/upload/',
            {'file': self.pdf()}, format='multipart')
        document = ResearchDocument.objects.get(pk=uploaded.data['id'])
        version = document.current_version
        version.processing_status = DocumentVersion.ProcessingStatus.PROCESSING
        version.processing_started_at = timezone.now()
        version.save(update_fields=('processing_status', 'processing_started_at'))
        retry_url = f'/api/v1/documents/{document.pk}/versions/{version.pk}/retry-extraction/'
        with patch('apps.research.views.process_document_version.delay') as enqueue:
            response = self.client.post(retry_url, {}, format='json')
        self.assertEqual(response.status_code, 409)
        enqueue.assert_not_called()
        self.client.force_authenticate(self.other)
        self.assertEqual(self.client.post(retry_url, {}, format='json').status_code, 404)

    def test_stale_processing_retry_fences_late_worker_result(self):
        project = self.project()
        uploaded = self.client.post(f'/api/v1/projects/{project.pk}/documents/upload/',
            {'file': self.pdf()}, format='multipart')
        document = ResearchDocument.objects.get(pk=uploaded.data['id'])
        version = document.current_version
        version.processing_status = DocumentVersion.ProcessingStatus.PROCESSING
        version.processing_started_at = timezone.now() - timedelta(minutes=20)
        version.processing_attempt = 1
        version.save(update_fields=('processing_status', 'processing_started_at', 'processing_attempt'))
        retry_url = f'/api/v1/documents/{document.pk}/versions/{version.pk}/retry-extraction/'
        with override_settings(DOCUMENT_PROCESSING_STALE_SECONDS=60), patch(
            'apps.research.views.process_document_version.delay') as enqueue:
            response = self.client.post(retry_url, {}, format='json')
        self.assertEqual(response.status_code, 202)
        enqueue.assert_called_once_with(str(version.pk))
        version.refresh_from_db()
        self.assertEqual(version.processing_status, DocumentVersion.ProcessingStatus.PENDING)

        with patch('apps.research.tasks.extract_document', return_value={'sections': [{'title': 'New result'}]}):
            process_document_version.run(str(version.pk))
        version.refresh_from_db()
        self.assertEqual(version.processing_attempt, 2)
        self.assertEqual(version.processing_status, DocumentVersion.ProcessingStatus.COMPLETED)
        self.assertIn('New result', str(version.extraction_metadata))

    def test_superseded_worker_cannot_overwrite_retry_result(self):
        project = self.project()
        uploaded = self.client.post(f'/api/v1/projects/{project.pk}/documents/upload/',
            {'file': self.pdf()}, format='multipart')
        document = ResearchDocument.objects.get(pk=uploaded.data['id'])
        version = document.current_version
        version.processing_status = DocumentVersion.ProcessingStatus.PENDING
        version.processing_started_at = None
        version.processing_attempt = 0
        version.save(update_fields=('processing_status', 'processing_started_at', 'processing_attempt'))
        def finish_after_retry(path, extension):
            # Simulate the operator reclaiming an expired worker's lease and a
            # newer task completing before the old parser returns.
            DocumentVersion.objects.filter(pk=version.pk).update(
                processing_status=DocumentVersion.ProcessingStatus.PENDING, processing_started_at=None)
            with patch('apps.research.tasks.extract_document', return_value={'result': 'new worker result'}):
                process_document_version.run(str(version.pk))
            return {'result': 'late stale worker result'}

        with patch('apps.research.tasks.extract_document', side_effect=finish_after_retry):
            old_result = process_document_version.run(str(version.pk))
        version.refresh_from_db()
        self.assertTrue(old_result['stale_attempt'])
        self.assertEqual(version.processing_status, DocumentVersion.ProcessingStatus.COMPLETED)
        self.assertEqual(version.extraction_metadata, {'result': 'new worker result'})
        self.assertEqual(version.processing_attempt, 2)
